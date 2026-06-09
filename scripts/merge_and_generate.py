"""
scripts/merge_and_generate.py
------------------------------
Reads all 6 observation xlsx files from tests/mevid_xlsx/,
cleans & normalises them, merges into one sightings list,
then generates agent_test_mevid.xlsx (Part1 sheet) ready for
agent_test_importer.py / ragas_eval_runner.py.

Slot → annotator mapping (derived from camera sets):
  11-20  →  Cz               (Person 1,  8 cameras)
  11-55  →  Shunyu Zhang     (Person 2,  7 cameras)
  13-50  →  MinqiangYao      (Person 3,  7 cameras)
  14-20  →  Quanjie Sheng    (Person 4,  8 cameras)
  16-20  →  Junwen Chen      (Person 5,  8 cameras)
  16-35  →  Jiaxin           (Person 6, 10 cameras)

Usage:
    python scripts/merge_and_generate.py
    python scripts/merge_and_generate.py --max-neg 40 --seed 0
    python scripts/merge_and_generate.py --out agent/test/data/agent_test.xlsx
"""

from __future__ import annotations

import argparse
import datetime
import random
import re
from collections import defaultdict
from itertools import combinations
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Paths ─────────────────────────────────────────────────────────────────────
ROOT = Path(__file__).resolve().parents[1]
XLSX_DIR = ROOT / "tests" / "mevid_xlsx"
DEFAULT_OUT = ROOT / "agent" / "test" / "data" / "agent_test_mevid.xlsx"

# ── File → slot assignment ─────────────────────────────────────────────────────
FILE_SLOT: dict[str, str] = {
    "observation_form_Cz.xlsx":                        "11-20",
    "Shunyu  observation_form(2).xlsx":                "11-55",
    "observation_form_from_MinqiangYao.xlsx":          "13-50",
    "observation_form_QuanjieSheng(1).xlsx":           "14-20",
    "observation_form_JunwenChen_integrated.xlsx":     "16-20",
    "副本observation_form.xlsx":                        "16-35",
}

# ── (slot, camera) → video stem (filename without .avi) ─────────────────────
# NOTE: end-times are approximate (+5 min); location codes for G299/G330/G421/
#       G424/G505 are estimated.  Verify against the MEVID URL list if needed.
SLOT_CAMERAS: dict[str, dict[str, str]] = {
    "11-20": {
        "G329": "2018-03-11.11-20-00.11-25-00.admin.G329.r13",
        "G330": "2018-03-11.11-20-00.11-25-00.school.G330.r13",
        "G336": "2018-03-11.11-20-00.11-25-00.school.G336.r13",
        "G419": "2018-03-11.11-20-00.11-25-00.school.G419.r13",
        "G420": "2018-03-11.11-20-00.11-25-00.school.G420.r13",
        "G421": "2018-03-11.11-20-00.11-25-00.school.G421.r13",
        "G423": "2018-03-11.11-20-00.11-25-00.school.G423.r13",
        "G508": "2018-03-11.11-20-00.11-25-00.bus.G508.r13",
    },
    "11-55": {
        "G299": "2018-03-11.11-55-00.12-00-00.school.G299.r13",
        "G328": "2018-03-11.11-55-00.12-00-00.school.G328.r13",
        "G330": "2018-03-11.11-55-00.12-00-00.school.G330.r13",
        "G419": "2018-03-11.11-55-00.12-00-00.school.G419.r13",
        "G420": "2018-03-11.11-55-00.12-00-00.school.G420.r13",
        "G506": "2018-03-11.11-55-00.12-00-00.bus.G506.r13",
        "G508": "2018-03-11.11-55-00.12-00-00.bus.G508.r13",
    },
    "13-50": {
        "G328": "2018-03-11.13-50-01.13-55-01.school.G328.r13",
        "G329": "2018-03-11.13-50-01.13-55-01.admin.G329.r13",
        "G339": "2018-03-11.13-50-01.13-55-01.school.G339.r13",
        "G421": "2018-03-11.13-50-01.13-55-01.school.G421.r13",
        "G424": "2018-03-11.13-50-01.13-55-01.school.G424.r13",
        "G506": "2018-03-11.13-50-01.13-55-01.bus.G506.r13",
        "G508": "2018-03-11.13-50-01.13-55-01.bus.G508.r13",
    },
    "14-20": {
        "G328": "2018-03-11.14-20-01.14-25-01.school.G328.r13",
        "G339": "2018-03-11.14-20-01.14-25-01.school.G339.r13",
        "G419": "2018-03-11.14-20-01.14-25-01.school.G419.r13",
        "G421": "2018-03-11.14-20-01.14-25-01.school.G421.r13",
        "G423": "2018-03-11.14-20-01.14-25-01.school.G423.r13",
        "G505": "2018-03-11.14-20-01.14-25-01.bus.G505.r13",
        "G506": "2018-03-11.14-20-01.14-25-01.bus.G506.r13",
        "G508": "2018-03-11.14-20-01.14-25-01.bus.G508.r13",
    },
    "16-20": {
        "G326": "2018-03-11.16-20-01.16-25-01.admin.G326.r13",
        "G328": "2018-03-11.16-20-01.16-25-01.school.G328.r13",
        "G329": "2018-03-11.16-20-01.16-25-01.admin.G329.r13",
        "G336": "2018-03-11.16-20-01.16-25-01.school.G336.r13",
        "G419": "2018-03-11.16-20-01.16-25-01.school.G419.r13",
        "G420": "2018-03-11.16-20-01.16-25-01.school.G420.r13",
        "G506": "2018-03-11.16-20-01.16-25-01.bus.G506.r13",
        "G508": "2018-03-11.16-20-01.16-25-01.bus.G508.r13",
    },
    "16-35": {
        "G326": "2018-03-11.16-35-01.16-40-01.admin.G326.r13",
        "G328": "2018-03-11.16-35-01.16-40-01.school.G328.r13",
        "G329": "2018-03-11.16-35-01.16-40-01.admin.G329.r13",
        "G336": "2018-03-11.16-35-01.16-40-01.school.G336.r13",
        "G339": "2018-03-11.16-35-01.16-40-01.school.G339.r13",
        "G419": "2018-03-11.16-35-01.16-40-01.school.G419.r13",
        "G420": "2018-03-11.16-35-01.16-40-01.school.G420.r13",
        "G423": "2018-03-11.16-35-01.16-40-00.school.G423.r13",
        "G506": "2018-03-11.16-35-01.16-40-01.bus.G506.r13",
        "G638": "2018-03-11.16-35-01.16-40-01.school.G638.r13",
    },
}


# ── Time normalisation ────────────────────────────────────────────────────────
def normalise_time(val) -> str:
    """Convert any time cell value to 'm:ss' string.

    Excel stores 'mm:ss' entries as time values.  openpyxl reads them as
    datetime.time where .hour = minutes and .minute = seconds (because the
    annotators entered e.g. '3:11' which Excel treats as h:mm = 3h 11m, but
    the semantic meaning in a 5-min clip is 3 min 11 sec).

    Cz used decimal notation: 0.18 = 0:18, 1.34 = 1:34 (M.SS).
    JunwenChen / Jiaxin used plain strings: '0:28'.
    """
    if val is None:
        return ""
    if isinstance(val, datetime.time):
        # .hour holds the minutes-part, .minute holds the seconds-part
        return f"{val.hour}:{val.minute:02d}"
    if isinstance(val, float):
        # Cz decimal format: 1.34 → 1:34
        minutes = int(val)
        seconds = round((val - minutes) * 100)
        if seconds >= 60:          # handle floating-point artefacts like 0.60
            minutes += seconds // 60
            seconds %= 60
        return f"{minutes}:{seconds:02d}"
    if isinstance(val, int):
        return f"{val}:00"
    # Already a string
    return str(val).strip()


# ── P-tag helpers ─────────────────────────────────────────────────────────────
_PTAG_RE = re.compile(r'^[Pp]\d+$')

def is_valid_ptag(tag: str) -> bool:
    return bool(_PTAG_RE.match(tag.strip()))

def normalise_ptag(tag: str) -> str:
    """Upper-case first letter: p1 → P1."""
    t = tag.strip()
    if t and t[0].lower() == 'p':
        return 'P' + t[1:]
    return t

def prefix_tag(slot: str, tag: str) -> str:
    return f"{slot}-{normalise_ptag(tag)}"

def parse_same_as(same_as_raw: str, slot: str) -> str | None:
    """Return the canonical prefixed tag, or None if same-camera or empty.

    Handles:
      - 'P10(same camera)'   → None   (re-entry in same camera, skip)
      - 'P10, P16'           → first valid tag prefixed with slot
      - 'P5,P18'             → first valid tag
      - 'p1'  (lower-case)   → normalised + prefixed
      - None / empty         → None
    """
    raw = (same_as_raw or "").strip()
    if not raw:
        return None
    raw_lower = raw.lower()
    if "same camera" in raw_lower or "same cam" in raw_lower:
        return None
    # Strip parenthetical annotations like "(same camera)" or "(G328)"
    raw = re.sub(r'\(.*?\)', '', raw).strip()
    # May be comma-separated list; take the first valid one
    for part in re.split(r'[,;/]', raw):
        part = part.strip()
        if is_valid_ptag(part):
            return prefix_tag(slot, part)
    return None


# ── Load one xlsx file ────────────────────────────────────────────────────────
_EXAMPLE_ANNOTATORS = {"alice"}   # rows from the template example; skip

def load_xlsx(path: Path, slot: str) -> list[dict]:
    """Load sightings from one observation xlsx, return normalised records."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sightings"] if "Sightings" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))

    # Find the header row
    header_idx = None
    for i, row in enumerate(rows):
        if any(str(v or "").strip().lower() == "person_tag" for v in row):
            header_idx = i
            break
    if header_idx is None:
        print(f"  [WARN] no header row in {path.name}, skipping")
        return []

    headers = [str(v or "").strip().lower() for v in rows[header_idx]]

    def get(row, col):
        try:
            idx = headers.index(col)
            val = row[idx] if idx < len(row) else None
            return val  # raw; caller normalises
        except ValueError:
            return None

    def gets(row, col) -> str:
        val = get(row, col)
        if val is None:
            return ""
        return str(val).strip()

    sightings: list[dict] = []
    for row in rows[header_idx + 1:]:
        if not any(v for v in row):
            continue

        raw_tag  = str(get(row, "person_tag") or "").strip()
        raw_cam  = gets(row, "camera_id").upper()
        annotator = gets(row, "annotator")

        # Skip template example rows
        if annotator.lower() in _EXAMPLE_ANNOTATORS:
            continue
        # Skip garbled header rows (e.g. '01', '02', '03' from JunwenChen)
        if not is_valid_ptag(raw_tag):
            continue
        if not raw_cam:
            continue
        # Only keep cameras that belong to this slot
        if raw_cam not in SLOT_CAMERAS.get(slot, {}):
            continue

        raw_same_as = gets(row, "same_as")
        canonical_same_as = parse_same_as(raw_same_as, slot)

        # Determine final person_tag:
        # If same_as points to another person in a DIFFERENT camera → rename to canonical
        own_prefixed = prefix_tag(slot, raw_tag)
        final_tag = canonical_same_as if canonical_same_as else own_prefixed

        sightings.append({
            "slot":           slot,
            "annotator":      annotator,
            "person_tag":     final_tag,
            "camera_id":      raw_cam,
            "time_enter":     normalise_time(get(row, "time_enter")),
            "time_exit":      normalise_time(get(row, "time_exit")),
            "exit_direction": gets(row, "exit_direction").lower(),
            "appearance":     gets(row, "appearance"),
            "same_as_raw":    raw_same_as,
            "confidence":     gets(row, "confidence") or "high",
            "notes":          gets(row, "notes"),
            "video_id":       SLOT_CAMERAS[slot].get(raw_cam, f"UNKNOWN-{slot}-{raw_cam}"),
        })

    return sightings


# ── Load all 6 files ──────────────────────────────────────────────────────────
def load_all(xlsx_dir: Path) -> list[dict]:
    all_sightings: list[dict] = []
    for fname, slot in FILE_SLOT.items():
        path = xlsx_dir / fname
        if not path.exists():
            print(f"[WARN] file not found: {path}")
            continue
        recs = load_xlsx(path, slot)
        print(f"  {fname[:45]:<45} slot={slot}  rows={len(recs)}")
        all_sightings.extend(recs)
    return all_sightings


# ── Query generation ──────────────────────────────────────────────────────────
def extract_top(appearance: str) -> str:
    if not appearance:
        return ""
    parts = [p.strip() for p in appearance.split(",") if p.strip()]
    return parts[0] if parts else appearance

def time_range_str(s: dict) -> str | None:
    te, tx = s["time_enter"], s["time_exit"]
    if te and tx and te != tx:
        return f"{te}-{tx}"
    return te or tx or None


def generate(sightings: list[dict], max_neg: int, seed: int) -> list[dict]:
    random.seed(seed)
    cases: list[dict] = []

    # Index by person_tag
    by_person: dict[str, list[dict]] = defaultdict(list)
    for s in sightings:
        by_person[s["person_tag"]].append(s)

    # Index confirmed cross-camera pairs (for negative filtering)
    confirmed_cam_pairs: set[tuple[str, str]] = set()
    for tag, obs in by_person.items():
        cams = [o["camera_id"] for o in obs]
        for ca, cb in combinations(cams, 2):
            if ca != cb:
                confirmed_cam_pairs.add(tuple(sorted([ca, cb])))  # type: ignore[arg-type]

    def make(video_id, question, challenge, answer, t, difficulty, category) -> dict:
        return {
            "video_id":        video_id,
            "question":        question,
            "recall_challenge": challenge,
            "expected_answer": answer,
            "expected_time":   t or "N/A",
            "difficulty":      difficulty,
            "video_quality":   "high",
            "_category":       category,
        }

    # 1. Existence  (easy, one per sighting)
    for s in sightings:
        top = extract_top(s["appearance"])
        q = (f"Is there a person with {top} visible in camera {s['camera_id']}?"
             if top else
             f"Is there a person visible in camera {s['camera_id']}?")
        cases.append(make(s["video_id"], q,
                          s["appearance"] or f"person in {s['camera_id']}",
                          "yes", time_range_str(s), "easy", "existence"))

    # 2. Appearance attribute (medium)
    for s in sightings:
        if not s["appearance"]:
            continue
        parts = [p.strip() for p in s["appearance"].split(",") if p.strip()]
        t = time_range_str(s)
        # Top clothing
        if parts:
            q = f"Is there a person wearing {parts[0]} in camera {s['camera_id']}?"
            cases.append(make(s["video_id"], q, s["appearance"], "yes", t, "medium", "appearance"))
        # Bag / backpack
        bag = next((p for p in parts if any(w in p.lower()
                    for w in ["bag", "backpack", "tote", "carry", "briefcase", "handbag"])), None)
        if bag:
            q = f"Is there a person carrying {bag} in camera {s['camera_id']}?"
            cases.append(make(s["video_id"], q, s["appearance"], "yes", t, "medium", "appearance"))

    # 3. Exit direction  (medium)
    for s in sightings:
        d = s["exit_direction"]
        if not d or d in ("stayed", "still_in_frame", ""):
            continue
        t = time_range_str(s)
        q = f"Did a person exit from the {d} side of camera {s['camera_id']}?"
        challenge = f"Person leaving frame {d}; {s['appearance']}"
        cases.append(make(s["video_id"], q, challenge, "yes", t, "medium", "event"))

    # 4. Cross-camera positive  (hard)
    for tag, obs in by_person.items():
        multi = obs  # all sightings for this canonical identity
        seen_pairs: set[tuple[str, str]] = set()
        for i, oa in enumerate(multi):
            for j, ob in enumerate(multi):
                if i >= j:
                    continue
                ca, cb = oa["camera_id"], ob["camera_id"]
                if ca == cb:
                    continue
                pair_key = tuple(sorted([ca, cb]))  # type: ignore[arg-type]
                if pair_key in seen_pairs:
                    continue
                seen_pairs.add(pair_key)  # type: ignore[arg-type]

                # Put earlier camera first
                def _sec(t: str) -> float:
                    try:
                        m, s = t.split(":")
                        return int(m) * 60 + int(s)
                    except Exception:
                        return 9999.0

                first, second = (
                    (oa, ob) if _sec(oa["time_enter"]) <= _sec(ob["time_enter"]) else (ob, oa)
                )
                top = extract_top(first["appearance"])
                conf = second["confidence"] or first["confidence"] or "high"
                diff = "medium" if conf == "medium" else "hard"
                t = time_range_str(second)
                desc = top if top else f"person {tag}"
                q = (f"Did a person with {desc} appear in camera {first['camera_id']}"
                     f" and then appear again in camera {second['camera_id']}?")
                challenge = (
                    f"Same person ({tag}) first in {first['camera_id']} @ "
                    f"{first['time_enter'] or '?'}, then {second['camera_id']} @ "
                    f"{second['time_enter'] or '?'}. "
                    f"Appearance: {first['appearance']}"
                )
                cases.append(make(second["video_id"], q, challenge, "yes", t, diff, "cross_camera"))

    # 5. Negative / FP  (medium)
    # Only generate negatives WITHIN the same slot (same time-window) to be meaningful.
    unique: dict[tuple, dict] = {}
    for s in sightings:
        k = (s["person_tag"], s["camera_id"])
        if k not in unique:
            unique[k] = s
    unique_list = list(unique.values())

    neg_candidates: list[tuple[dict, dict]] = []
    for i, sa in enumerate(unique_list):
        for sb in unique_list[i + 1:]:
            if sa["slot"] != sb["slot"]:          # must be same slot
                continue
            if sa["camera_id"] == sb["camera_id"]:
                continue
            if sa["person_tag"] == sb["person_tag"]:
                continue
            pair = tuple(sorted([sa["camera_id"], sb["camera_id"]]))  # type: ignore[arg-type]
            if pair in confirmed_cam_pairs:       # already a confirmed cross-cam match
                continue
            neg_candidates.append((sa, sb))

    random.shuffle(neg_candidates)
    for sa, sb in neg_candidates[:max_neg]:
        top_a = extract_top(sa["appearance"])
        desc = top_a if top_a else f"person {sa['person_tag']}"
        q = (f"Did a person with {desc} from camera {sa['camera_id']}"
             f" also appear in camera {sb['camera_id']}?")
        challenge = (
            f"NOT a match. {sa['person_tag']} ({sa['appearance']}) in {sa['camera_id']} "
            f"≠ {sb['person_tag']} ({sb['appearance']}) in {sb['camera_id']}"
        )
        cases.append(make(sa["video_id"], q, challenge, "no", "N/A", "medium", "negative"))

    return cases


# ── Write agent_test xlsx ─────────────────────────────────────────────────────
CAT_COLOUR = {
    "existence":   "DEEAF1",
    "appearance":  "E2EFDA",
    "event":       "FFF2CC",
    "cross_camera": "F4CCFF",
    "negative":    "FFE0E0",
}
_thin = Side(style="thin", color="CCCCCC")


def _fill(h: str) -> PatternFill:
    return PatternFill("solid", fgColor=h)

def _bdr() -> Border:
    return Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

def _fnt(bold=False, color="000000", size=10) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size)


def write_xlsx(cases: list[dict], out_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Part1"

    headers = ["video_id", "question", "recall_challenge",
               "expected_answer", "expected_time", "difficulty", "video_quality"]
    widths   = [56, 84, 72, 16, 18, 12, 14]

    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        c.fill      = _fill("1F4E79")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _bdr()
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 24

    prev_vid = None
    for ri, case in enumerate(cases, 2):
        vid = case["video_id"]
        display_vid = vid if vid != prev_vid else None
        prev_vid = vid

        cat = case.get("_category", "")
        row_bg = CAT_COLOUR.get(cat, "FFFFFF")
        ans = case["expected_answer"]

        values = [display_vid, case["question"], case["recall_challenge"],
                  ans, case["expected_time"], case["difficulty"], case["video_quality"]]
        for ci, val in enumerate(values, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            if ci == 4:
                c.fill = _fill("C6EFCE" if ans == "yes" else "FFDCE1" if ans == "no" else row_bg)
                c.font = Font(name="Calibri", bold=True, size=10,
                              color="375623" if ans == "yes" else
                                    "C00000" if ans == "no" else "000000")
            else:
                c.fill = _fill(row_bg)
                c.font = _fnt(size=10)
            c.alignment = Alignment(vertical="center", wrap_text=True,
                                    horizontal="center" if ci in (4, 5, 6, 7) else "left")
            c.border = _bdr()
        ws.row_dimensions[ri].height = 38

    ws.freeze_panes = "A2"
    ws.sheet_view.zoomScale = 95

    # Legend sheet
    wl = wb.create_sheet("Legend")
    wl.column_dimensions["A"].width = 16
    wl.column_dimensions["B"].width = 65
    desc_map = {
        "existence":   "Is there a person with [desc] in camera [X]?",
        "appearance":  "Is there a person wearing/carrying [attr] in camera [X]?",
        "event":       "Did a person exit from the [direction] side of camera [X]?",
        "cross_camera": "Did [person] appear in [camA] and then in [camB]?  (answer=yes)",
        "negative":    "Cross-camera pair NOT confirmed as same person  (answer=no)",
    }
    for ri, (cat, bg) in enumerate(CAT_COLOUR.items(), 1):
        c1 = wl.cell(row=ri, column=1, value=cat)
        c1.fill  = _fill(bg); c1.font = _fnt(bold=True); c1.border = _bdr()
        c2 = wl.cell(row=ri, column=2, value=desc_map.get(cat, ""))
        c2.font  = _fnt(); c2.border = _bdr()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ── Main ──────────────────────────────────────────────────────────────────────
def get_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Merge 6 observation xlsx files and generate MEVID QA test cases")
    p.add_argument("--xlsx-dir", default=str(XLSX_DIR))
    p.add_argument("--out",      default=str(DEFAULT_OUT),
                   help="Output xlsx path (default: agent/test/data/agent_test_mevid.xlsx)")
    p.add_argument("--max-neg",  type=int, default=40,
                   help="Max negative (FP) cross-camera cases to generate (default: 40)")
    p.add_argument("--seed",     type=int, default=42)
    return p.parse_args()


def main() -> None:
    args = get_args()
    xlsx_dir = Path(args.xlsx_dir)
    out_path = Path(args.out)

    print("Loading observation files …")
    sightings = load_all(xlsx_dir)
    print(f"\nTotal sightings loaded: {len(sightings)}")

    if not sightings:
        print("ERROR: no sightings found.")
        return

    # Stats
    from collections import Counter
    slot_counts = Counter(s["slot"] for s in sightings)
    cam_counts  = Counter(s["camera_id"] for s in sightings)
    print("\nSightings per slot:", dict(sorted(slot_counts.items())))
    print("Unique person_tags:", len({s["person_tag"] for s in sightings}))

    print(f"\nGenerating test cases (seed={args.seed}, max_neg={args.max_neg}) …")
    cases = generate(sightings, args.max_neg, args.seed)

    cat_counts = Counter(c["_category"] for c in cases)
    ans_counts = Counter(c["expected_answer"] for c in cases)
    diff_counts = Counter(c["difficulty"] for c in cases)
    print(f"Generated {len(cases)} cases")
    print("  By category:", dict(cat_counts))
    print("  By answer:  ", dict(ans_counts))
    print("  By difficulty:", dict(diff_counts))

    write_xlsx(cases, out_path)
    print(f"\nSaved: {out_path}")
    print("\nNext steps:")
    print(f"  1. python agent/test/agent_test_importer.py --xlsx-path {out_path}")
    print(f"  2. python scripts/download_mevid_slots.py")
    print(f"  3. python tests/test_mevid_qa.py --xlsx {out_path}")


if __name__ == "__main__":
    main()
